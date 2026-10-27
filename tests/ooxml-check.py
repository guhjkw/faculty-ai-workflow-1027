"""tests/ooxml-check.py: opens the files tests/verify.js wrote (tests/out/) with python-pptx and python-docx,
independent readers that know nothing about our writer, and checks what they find.
Run:  pip install python-pptx python-docx && node tests/verify.js && python3 tests/ooxml-check.py"""
import pathlib, sys
from pptx import Presentation
from docx import Document

OUT = pathlib.Path(__file__).resolve().parent / 'out'
passed = failed = 0
def ok(name, cond, detail=''):
    global passed, failed
    passed, failed = (passed + 1, failed) if cond else (passed, failed + 1)
    print(('  PASS ' if cond else '  FAIL ') + name + ('' if cond else '  -> ' + str(detail)))

print('\n1. Instructor deck (final) in python-pptx')
p = Presentation(OUT / 'instructor-final.pptx')
slides = list(p.slides)
ok('widescreen 16:9 (12192000 x 6858000 EMU)', (p.slide_width, p.slide_height) == (12192000, 6858000))
ok('23 slides, the first on the Title Slide layout', len(slides) == 23 and slides[0].slide_layout.name == 'Title Slide', len(slides))
ok('every content slide uses Title and Content with a real title placeholder', all(s.slide_layout.name == 'Title and Content' and s.shapes.title is not None for s in slides[1:]))
key = [s for s in slides if s.shapes.title.text == 'Answer key'][0]
ok('the answer key is a hidden slide', key._element.get('show') == '0')
ok('and carries the INSTRUCTOR ONLY tag', any(sh.name == 'Instructor tag' and sh.text_frame.text == 'INSTRUCTOR ONLY' for sh in key.shapes))
cost = [s for s in slides if s.shapes.title.text == 'Cost Leadership'][0]
notes = cost.notes_slide.notes_text_frame.text
ok('speaker notes hold the instructor-only content', 'Probe: If your costs are the lowest' in notes and 'Talking point (publisher notes)' in notes, notes[:200])
body = [sh for sh in cost.placeholders if sh.placeholder_format.idx == 1][0]
ok('the body placeholder holds the 5 student-facing bullets', len([para for para in body.text_frame.paragraphs if para.text.strip()]) == 5)
ok('the source line says where the slide came from and that it is approved', any(sh.name == 'Source line' and 'Publisher slide 6' in sh.text_frame.text and 'approved' in sh.text_frame.text for sh in cost.shapes))

print('\n2. Student deck (draft)')
s = Presentation(OUT / 'student-draft.pptx')
ok('21 slides and no notes on any of them', len(s.slides) == 21 and not any(sl.has_notes_slide for sl in s.slides), len(s.slides))
ok('no answer key in the student copy', not any(sl.shapes.title.text == 'Answer key' for sl in s.slides))
ok('every slide is marked DRAFT', all(any(sh.name == 'Source line' and sh.text_frame.text.startswith('DRAFT') for sh in sl.shapes) for sl in list(s.slides)[1:]))

print('\n3. Word files in python-docx')
g = Document(OUT / 'instructor-guide.docx')
styles = [para.style.name for para in g.paragraphs]
ok('the guide opens with a Title and uses real Heading styles', styles[0] == 'Title' and styles.count('Heading 1') >= 6, styles[:6])
ok('its timing table adds to 75 minutes', sum(int(r.cells[1].text) for r in g.tables[0].rows[1:]) == 75)
ok('its source record has a row per slide plus a header', len(g.tables[1].rows) == 24, len(g.tables[1].rows))
c = Document(OUT / 'student-case.docx')
text = '\n'.join(para.text for para in c.paragraphs)
ok('the student case is disguised and has no epilogue', 'Pelican Point Marine' in text and 'Bayshore' not in text and '7.4 million' not in text)
ok('its questions are a real numbered list', sum(1 for para in c.paragraphs if para.style.name == 'List Number') == 3)
t = Document(OUT / 'teaching-note.docx')
ok('the teaching note carries the 8-fact master record', len(t.tables[0].rows) == 9)
m = Document(OUT / 'research-memo.docx')
ok('the research memo lists two verified references and no fake', sum(1 for para in m.paragraphs if para.text.startswith(('Barney', 'Porter'))) == 2 and not any('Whitlow' in para.text for para in m.paragraphs))

print('\n%d file checks: %d passed, %d failed' % (passed + failed, passed, failed))
sys.exit(1 if failed else 0)
