"""Writes tests/fixtures/publisher-short.pptx with python-pptx (a different writer from ours, compressed like
PowerPoint's own files), so the browser test can load a real publisher-style deck. Run once:
python3 tests/make-fixture.py   (needs python-pptx)"""
import pathlib
from pptx import Presentation
slides = [
    ('Chapter 6: Business-Level Strategy (short version)', []),
    ('What Is Business-Level Strategy?', ['An integrated and coordinated set of commitments and actions', 'Answers how the firm will compete']),
    ('Five Generic Strategies', ['Cost leadership', 'Differentiation', 'Focused cost leadership', 'Focused differentiation', 'Integrated cost leadership and differentiation']),
    ('Cost Leadership', ['Produce at the lowest cost relative to rivals', 'Standardized products for the typical customer', 'Value chain activities: lean operations & low overhead']),
    ('Differentiation', ['Customers pay a premium price for unique features', 'Sources include quality, brand image and service']),
    ('Risks of Each Strategy', ['Cost leadership: rivals imitate', 'Differentiation: the premium grows too large', 'Integrated: the firm can become stuck in the middle']),
]
p = Presentation()
for i, (title, bullets) in enumerate(slides):
    s = p.slides.add_slide(p.slide_layouts[0 if not bullets else 1])
    s.shapes.title.text = title
    if bullets:
        tf = s.placeholders[1].text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
        s.notes_slide.notes_text_frame.text = 'Publisher note for slide %d' % (i + 1)
out = pathlib.Path(__file__).resolve().parent / 'fixtures' / 'publisher-short.pptx'
p.save(out)
print(out, out.stat().st_size, 'bytes')
